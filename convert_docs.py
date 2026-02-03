import os
from docx import Document
from fpdf import FPDF
import logging

# Set up logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
logger = logging.getLogger(__name__)

__all__ = ['create_pdf']

def create_pdf(mission_name, content, output_dir):
    """Creates a PDF file for a mission."""
    try:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)
        
        # Add title
        pdf.set_font('Arial', 'B', 16)
        title = mission_name.replace('_', ' ').title()
        pdf.cell(0, 10, title, ln=True, align='C')
        pdf.ln(10)
        
        # Add content
        pdf.set_font('Arial', '', 12)
        for line in content:
            # Split long lines to fit page width
            words = line.split()
            line_buf = ''
            for word in words:
                test_line = line_buf + word + ' '
                if pdf.get_string_width(test_line) > 180:
                    pdf.multi_cell(0, 10, line_buf.strip())
                    line_buf = word + ' '
                else:
                    line_buf = test_line
            if line_buf:
                pdf.multi_cell(0, 10, line_buf.strip())
            pdf.ln(5)
        
        # Save the PDF
        output_path = os.path.join(output_dir, f'{mission_name.lower()}.pdf')
        pdf.output(output_path)
        logger.info(f"Created PDF: {output_path}")
        return True
    except Exception as e:
        logger.error(f"Error creating PDF for {mission_name}: {str(e)}")
        return False

def convert_docx_to_pdfs():
    """Converts the DOCX file to individual PDFs for each mission."""
    logger.info("Starting document conversion process")
    try:
        base_dir = os.path.dirname(os.path.abspath(__file__))
        docx_path = os.path.join(base_dir, 'static', 'pdfs', 'FUTURE MISSIONS.pdf.docx')
        output_dir = os.path.join(base_dir, 'static', 'pdfs', 'missions')
        
        # Check if source document exists
        if not os.path.exists(docx_path):
            logger.error(f"Source document not found at {docx_path}")
            return False
        
        # Create output directory if needed
        if not os.path.exists(output_dir):
            os.makedirs(output_dir)
            logger.info(f"Created output directory: {output_dir}")
        
        # Load the document
        doc = Document(docx_path)
        
        current_mission = None
        mission_content = []
        success_count = 0
        
        # Process each paragraph
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
                
            # Check if this is a mission header
            if text.isupper() and len(text) > 5:
                # Save previous mission if exists
                if current_mission and mission_content:
                    if create_pdf(current_mission, mission_content, output_dir):
                        success_count += 1
                    mission_content = []
                    
                current_mission = text.replace(' ', '_')
            else:
                mission_content.append(text)
        
        # Save the last mission
        if current_mission and mission_content:
            if create_pdf(current_mission, mission_content, output_dir):
                success_count += 1
        
        logger.info(f"Conversion complete: {success_count} PDFs created successfully")
        return True
        
    except Exception as e:
        logger.error(f"Error during conversion: {str(e)}")
        return False

if __name__ == '__main__':
    logger.info("Starting document conversion...")
    if convert_docx_to_pdfs():
        logger.info("Document conversion completed successfully")
    else:
        logger.error("Document conversion failed")
